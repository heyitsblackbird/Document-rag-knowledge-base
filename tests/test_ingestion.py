import pytest
from docx import Document as DocxDocument

from services import ingestion_service
from services.ingestion_service import ingest_document, _cleaned_data


def test_cleaned_data_collapses_whitespace_and_nulls():
    assert _cleaned_data("hello \x00 \n\n  world  ") == "hello world"


def test_cleaned_data_handles_empty_string():
    assert _cleaned_data("") == ""


def test_ingest_document_raises_for_missing_file():
    with pytest.raises(FileNotFoundError):
        ingest_document("does/not/exist.txt")


def test_ingest_document_raises_for_unsupported_type(tmp_path):
    bad_file = tmp_path / "notes.exe"
    bad_file.write_text("binary junk")

    with pytest.raises(ValueError):
        ingest_document(str(bad_file))


def test_ingest_document_txt_is_a_single_untitled_page(tmp_path):
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("  Habits   compound\nover time.  ")

    doc = ingest_document(str(txt_file))

    assert doc["source"] == "notes.txt"
    assert doc["file-type"] == ".txt"
    assert doc["pages"] == [{"page": None, "text": "Habits compound over time."}]


def test_ingest_document_legacy_doc_is_rejected(tmp_path):
    doc_file = tmp_path / "notes.doc"
    doc_file.write_bytes(b"legacy binary word junk")

    with pytest.raises(ValueError):
        ingest_document(str(doc_file))


def test_ingest_document_docx_extracts_paragraph_text(tmp_path):
    docx_file = tmp_path / "notes.docx"
    document = DocxDocument()
    document.add_paragraph("Habits compound over time.")
    document.add_paragraph("Small changes add up.")
    document.save(str(docx_file))

    doc = ingest_document(str(docx_file))

    assert doc["source"] == "notes.docx"
    assert doc["file-type"] == ".docx"
    assert doc["pages"] == [{"page": None, "text": "Habits compound over time. Small changes add up."}]


def test_ingest_document_docx_with_no_text_returns_no_pages(tmp_path):
    docx_file = tmp_path / "blank.docx"
    DocxDocument().save(str(docx_file))

    doc = ingest_document(str(docx_file))

    assert doc["pages"] == []


class _FakePage:
    def __init__(self, text):
        self._text = text

    def extract_text(self):
        return self._text


class _FakeReader:
    def __init__(self, _path):
        self.pages = [_FakePage("Page one content."), _FakePage(""), _FakePage("Page three content.")]


def test_ingest_document_pdf_preserves_page_numbers_and_skips_blank_pages(tmp_path, monkeypatch):
    monkeypatch.setattr(ingestion_service, "PdfReader", _FakeReader)

    pdf_file = tmp_path / "book.pdf"
    pdf_file.write_bytes(b"%PDF-fake")

    doc = ingest_document(str(pdf_file))

    assert doc["pages"] == [
        {"page": 1, "text": "Page one content."},
        {"page": 3, "text": "Page three content."},
    ]
